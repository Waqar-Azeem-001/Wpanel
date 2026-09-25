"""CSV / XLSX / PDF exports, the report pages and the REST API."""
import csv
import re
from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO

import pytest
from django.utils import timezone

from apps.audit.models import AuditEvent
from apps.core.exceptions import ServiceError
from apps.reports import exports, services
from apps.reports.definitions import Column, Result

from .conftest import invoice_for, make_order, new_client, pay_on

pytestmark = pytest.mark.django_db

D = Decimal
TODAY = timezone.localdate()


@pytest.fixture
def sample():
    return Result(
        title="Sample", params={"From": "2026-09-01", "To": "2026-09-25"},
        columns=[Column("name", "Name"), Column("amount", "Amount", "money"), Column("count", "Count", "int"),
                 Column("when", "When", "datetime"), Column("day", "Day", "date")],
        rows=[{"name": "Ada", "amount": D("1234.50"), "count": 3, "when": timezone.now(), "day": TODAY},
              {"name": "=HYPERLINK(\"http://evil.example\",\"click\")", "amount": D("-5"), "count": 0, "when": None,
               "day": None},
              {"name": "+1+1", "amount": D("0"), "count": 1, "when": None, "day": None}],
        summary=[("Total", "money", D("1229.50")), ("Rows", "int", 3)], notes=["A note."])


def pages(pdf):
    return len(re.findall(rb"/Type /Page(?![s\w])", pdf))


# --- Files -----------------------------------------------------------------------------------------------------------

def test_csv_has_real_values_and_a_byte_order_mark(sample):
    data = exports.to_csv(sample)
    assert data.startswith(b"\xef\xbb\xbf")  # so Excel reads UTF-8
    rows = list(csv.reader(StringIO(data.decode("utf-8-sig"))))
    assert rows[0] == ["Name", "Amount", "Count", "When", "Day"]
    assert rows[1][0] == "Ada" and rows[1][1] == "1234.50" and rows[1][2] == "3" and rows[1][4] == TODAY.isoformat()
    assert rows[2][1] == "-5.00" and rows[2][3] == "" and len(rows) == 4  # no summary lines: it is pure data


def test_csv_neutralises_formulas(sample):
    rows = list(csv.reader(StringIO(exports.to_csv(sample).decode("utf-8-sig"))))
    assert rows[2][0].startswith("'=HYPERLINK") and rows[3][0] == "'+1+1"
    assert exports.to_csv(Result("T", [Column("x", "=SUM(A1)")], [])).decode("utf-8-sig").startswith("'=SUM(A1)")


def test_xlsx_keeps_numbers_dates_and_text_apart(sample):
    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(exports.to_xlsx(sample)))["Report"]
    cells = {row[0].value: row for row in sheet.iter_rows() if row and row[0].value}
    assert "Sample" in cells and cells["From"][1].value == "2026-09-01"
    assert cells["Total"][1].value == 1229.5 and cells["Rows"][1].value == 3  # figures are numbers
    header = cells["Name"]
    assert [c.value for c in header[:5]] == ["Name", "Amount", "Count", "When", "Day"] and header[0].font.bold
    first = cells["Ada"]
    assert first[1].value == 1234.5 and first[1].number_format == "#,##0.00" and first[1].data_type == "n"
    assert first[2].value == 3 and first[4].number_format == "yyyy-mm-dd"
    assert first[3].value is not None and first[3].number_format == "yyyy-mm-dd hh:mm"
    assert "A note." in cells


def test_xlsx_never_stores_a_formula(sample):
    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(exports.to_xlsx(sample)))["Report"]
    values = [c for row in sheet.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert values and all(c.data_type == "s" and c.value.startswith("'=") for c in values)
    assert not any(c.data_type == "f" for row in sheet.iter_rows() for c in row)


def test_pdf_is_a_real_document_that_paginates_and_is_repeatable(sample):
    pdf = exports.to_pdf(sample)
    assert pdf.startswith(b"%PDF") and pdf.rstrip().endswith(b"%%EOF")
    assert exports.to_pdf(sample) == pdf  # deterministic
    long = Result("Long", [Column("a", "A"), Column("b", "B", "money")],
                  [{"a": f"row {i}", "b": D(i)} for i in range(400)])
    assert pages(exports.to_pdf(sample)) == 1 and pages(exports.to_pdf(long)) > 1


def test_pdf_survives_characters_the_builtin_fonts_lack():
    result = Result("Ünïcode ✓ 漢", [Column("a", "Name")], [{"a": "Łukasz <b>&</b> 漢字"}])
    assert exports.to_pdf(result).startswith(b"%PDF")


def test_an_empty_report_still_exports(sample):
    empty = Result("Empty", [Column("a", "A")], [])
    for fmt in ("csv", "xlsx", "pdf"):
        assert exports.render(empty, fmt)[0]
    with pytest.raises(ValueError):
        exports.render(empty, "docx")


def test_exporting_is_audited_and_named_by_report_and_day(manager):
    new_client(manager, "Ada", "ada@example.com")
    content, content_type, filename = services.export(manager, "new-customers", {}, "csv")
    assert content_type.startswith("text/csv") and filename == f"new-customers-{TODAY:%Y%m%d}.csv"
    event = AuditEvent.objects.get(action="report.exported")
    assert event.actor == manager and event.metadata["report"] == "new-customers" and event.metadata["rows"] == 1
    assert event.metadata["format"] == "csv" and "From" in event.metadata["parameters"]


def test_export_rules(manager, agent):
    with pytest.raises(ServiceError) as exc:
        services.export(manager, "orders", {}, "docx")
    assert exc.value.code == "format_invalid"
    with pytest.raises(ServiceError) as exc:
        services.export(agent, "orders", {}, "csv")
    assert exc.value.code == "permission_denied" and not AuditEvent.objects.filter(action="report.exported").exists()


# --- Pages -----------------------------------------------------------------------------------------------------------

def test_the_index_lists_what_the_person_may_run(client, manager, agent, customer, support_reporter):
    pay_on(manager, invoice_for(manager, new_client(manager, "Ada", "ada@example.com"), "120.00"), 0)
    client.force_login(manager)
    page = client.get("/staff/reports/").content
    for title in (b"Sales performance", b"Income", b"Unpaid invoices", b"Active services", b"Open tickets"):
        assert title in page
    assert b"Income today (USD)" in page and b"120.00" in page and b"/staff/reports/" in client.get("/account/profile/").content
    client.force_login(support_reporter)
    page = client.get("/staff/reports/").content
    assert b"Open tickets" in page and b"Income" not in page and b"Unpaid invoices" not in page
    for user in (agent, customer):
        client.force_login(user)
        assert client.get("/staff/reports/").status_code == 403
    assert b"/staff/reports/" not in client.get("/account/profile/").content
    client.logout()
    assert client.get("/staff/reports/")["Location"].startswith("/account/login/")


def test_a_report_page_shows_the_table_the_filters_and_the_download_links(client, manager, bank):
    ada = new_client(manager, "Ada", "ada@example.com")
    pay_on(manager, invoice_for(manager, ada, "250.00"), 2, reference="TT-77", method=bank)
    client.force_login(manager)
    page = client.get(f"/staff/reports/income/?from={TODAY - timedelta(days=4)}&to={TODAY}&group=day").content
    assert b"250.00" in page and b'name="from"' in page and b'name="group"' in page and b"Received" in page
    assert b"export=csv" in page and b"export=xlsx" in page and b"export=pdf" in page
    assert f"from={TODAY - timedelta(days=4)}".encode() in page  # the links repeat what is on screen
    assert b"TT-77" in client.get("/staff/reports/transactions/").content
    assert b'name="days"' in client.get("/staff/reports/expiring-services/").content
    assert b'name="from"' not in client.get("/staff/reports/open-tickets/").content  # it takes no dates


def test_bad_parameters_are_explained_not_crashed(client, manager):
    client.force_login(manager)
    page = client.get("/staff/reports/income/?from=notadate&to=2026-09-01")
    assert page.status_code == 200 and b"must be a date" in page.content
    assert b"must not be after" in client.get("/staff/reports/orders/?from=2026-09-10&to=2026-09-01").content
    assert b"group by month" in client.get("/staff/reports/income/?from=2020-01-01&to=2026-01-01&group=day").content
    assert b"must be between 1 and 365" in client.get("/staff/reports/expiring-services/?days=999").content
    assert b"must be a date" in client.get("/staff/reports/orders/?from=x&export=csv").content  # a download too


def test_report_page_access(client, manager, agent, support_reporter):
    client.force_login(manager)
    assert client.get("/staff/reports/no-such-report/").status_code == 404
    client.force_login(agent)
    assert client.get("/staff/reports/orders/").status_code == 403
    client.force_login(support_reporter)
    assert client.get("/staff/reports/open-tickets/").status_code == 200
    assert client.get("/staff/reports/income/").status_code == 403
    assert client.get("/staff/reports/income/?export=csv").status_code == 403
    client.logout()
    assert client.get("/staff/reports/income/")["Location"].startswith("/account/login/")


@pytest.mark.parametrize("fmt,prefix,content_type", [
    ("csv", b"\xef\xbb\xbf", "text/csv"), ("xlsx", b"PK", "spreadsheetml"), ("pdf", b"%PDF", "application/pdf")])
def test_downloads_are_attachments_that_cannot_be_sniffed_or_cached(client, manager, fmt, prefix, content_type):
    make_order(new_client(manager, "Ada", "ada@example.com"))
    client.force_login(manager)
    response = client.get(f"/staff/reports/orders/?export={fmt}")
    assert response.status_code == 200 and response.content.startswith(prefix) and content_type in response["Content-Type"]
    assert response["Content-Disposition"] == f'attachment; filename="orders-{TODAY:%Y%m%d}.{fmt}"'
    assert response["X-Content-Type-Options"] == "nosniff" and "no-store" in response["Cache-Control"]
    assert AuditEvent.objects.filter(action="report.exported", metadata__format=fmt).count() == 1


def test_a_long_report_is_cut_on_the_page_but_complete_in_the_download(client, manager, monkeypatch):
    client_obj = new_client(manager, "Ada", "ada@example.com")
    for _ in range(5):
        make_order(client_obj)
    monkeypatch.setattr(services, "SCREEN_ROWS", 2)
    client.force_login(manager)
    page = client.get("/staff/reports/orders/").content
    assert b"Showing the first 2 of 5 rows" in page
    rows = list(csv.reader(StringIO(client.get("/staff/reports/orders/?export=csv").content.decode("utf-8-sig"))))
    assert len(rows) == 6  # the header and every order


def test_a_customers_name_cannot_run_as_a_formula_in_the_download(client, manager):
    new_client(manager, "=cmd|' /C calc'!A0", "evil@example.com")
    client.force_login(manager)
    body = client.get("/staff/reports/new-customers/?export=csv").content.decode("utf-8-sig")
    assert "'=cmd" in body and "\n=cmd" not in body and ",=cmd" not in body


# --- API -------------------------------------------------------------------------------------------------------------

def test_api_lists_only_the_reports_the_person_may_run(api, manager, agent, customer, support_reporter):
    api.force_authenticate(manager)
    listing = api.get("/api/v1/reports/").json()
    assert len(listing) == 14 and {r["group"] for r in listing} == {"Sales", "Financial", "Services", "Support"}
    income = next(r for r in listing if r["slug"] == "income")
    assert income["parameters"] == ["from", "to", "group"]
    assert next(r for r in listing if r["slug"] == "expiring-services")["parameters"] == ["days", "scope"]
    api.force_authenticate(support_reporter)
    assert {r["slug"] for r in api.get("/api/v1/reports/").json()} == {"open-tickets", "resolved-tickets",
                                                                       "department-stats"}
    for user in (agent, customer):
        api.force_authenticate(user)
        assert api.get("/api/v1/reports/").json() == []
    api.force_authenticate(None)
    assert api.get("/api/v1/reports/").status_code == 401


def test_api_returns_a_report_as_json(api, manager):
    pay_on(manager, invoice_for(manager, new_client(manager, "Ada", "ada@example.com"), "80.00"), 1)
    api.force_authenticate(manager)
    data = api.get("/api/v1/reports/income/", {"from": str(TODAY - timedelta(days=2)), "to": str(TODAY)}).json()
    assert data["slug"] == "income" and data["parameters"]["From"] == str(TODAY - timedelta(days=2))
    assert [c["key"] for c in data["columns"]] == ["period", "currency", "payments", "paid", "refunded", "net"]
    assert data["columns"][3]["kind"] == "money" and len(data["rows"]) == 3
    yesterday = next(r for r in data["rows"] if r["period"] == str(TODAY - timedelta(days=1)))
    assert yesterday["paid"] == "80.00" and yesterday["payments"] == 1  # money as exact strings, never floats
    assert {"label": "Net income", "kind": "money", "value": "80.00"} in data["summary"]
    assert data["truncated"] is False and data["notes"]


def test_api_errors_and_access(api, manager, support_reporter, agent):
    api.force_authenticate(manager)
    assert api.get("/api/v1/reports/nope/").status_code == 404
    bad = api.get("/api/v1/reports/income/", {"from": "x"})
    assert bad.status_code == 400
    assert api.get("/api/v1/reports/orders/", {"export": "docx"}).json()["error"]["code"] == "format_invalid"
    api.force_authenticate(support_reporter)
    assert api.get("/api/v1/reports/open-tickets/").status_code == 200
    assert api.get("/api/v1/reports/income/").status_code == 403
    assert api.get("/api/v1/reports/income/", {"export": "csv"}).status_code == 403
    api.force_authenticate(agent)
    assert api.get("/api/v1/reports/orders/").status_code == 403
    api.force_authenticate(None)
    assert api.get("/api/v1/reports/orders/").status_code == 401


@pytest.mark.parametrize("fmt,prefix", [("csv", b"\xef\xbb\xbf"), ("xlsx", b"PK"), ("pdf", b"%PDF")])
def test_api_downloads_files(api, manager, fmt, prefix):
    make_order(new_client(manager, "Ada", "ada@example.com"))
    api.force_authenticate(manager)
    response = api.get("/api/v1/reports/orders/", {"export": fmt})
    assert response.status_code == 200 and response.content.startswith(prefix)
    assert "attachment" in response["Content-Disposition"] and response["X-Content-Type-Options"] == "nosniff"
    assert AuditEvent.objects.filter(action="report.exported", metadata__format=fmt).exists()
